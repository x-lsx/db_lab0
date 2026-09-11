import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def generate_excel(json_path: str, output_excel: str):
    if not os.path.exists(json_path):
        print(f"[EXCEL ERROR] Файл с данными {json_path} не найден!")
        sys.exit(1)

    # Читаем данные из JSON
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    if not data:
        print("[EXCEL ERROR] Набор данных пуст.")
        sys.exit(1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка по студиям"

    headers = list(data[0].keys())
    
    ws.append(headers)
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = 25

    for row_data in data:
        ws.append([row_data[key] for key in headers])

    # Сохраняем итоговый файл
    wb.save(output_excel)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(1)
        
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    
    generate_excel(input_file, output_file)